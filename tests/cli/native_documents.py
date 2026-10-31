#!/usr/bin/env python3
"""Actual log -> XLSX -> event -> PPTX, with replaceable example-owned adapters."""
import argparse
import hashlib
import importlib.util
import json
from pathlib import Path
import platform
import shutil
import signal
import subprocess
import time

EXAMPLE = Path(__file__).resolve().parents[2] / "examples/log-report"


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--loom", type=Path, required=True)
    parser.add_argument("--host", type=Path, required=True)
    parser.add_argument("--python", type=Path, required=True)
    parser.add_argument("--host-revision", required=True)
    parser.add_argument("--output", type=Path, required=True)
    options = parser.parse_args()
    if platform.system() != "Darwin":
        raise SystemExit("Real macOS native enforcement required; no fallback")
    root = options.output.resolve()
    root.mkdir(mode=0o700)
    host, loom = options.host.resolve(strict=True), options.loom.resolve(strict=True)
    python = options.python.absolute()
    calls = []
    started = time.monotonic()

    def call(binary, *args):
        began = time.monotonic()
        process = subprocess.Popen([str(binary), *map(str, args)], text=True,
                                   stdout=subprocess.PIPE, stderr=subprocess.PIPE)
        try:
            stdout, stderr = process.communicate(timeout=60)
        except subprocess.TimeoutExpired as error:
            process.send_signal(signal.SIGINT)
            try:
                process.communicate(timeout=10)
            except subprocess.TimeoutExpired:
                raise RuntimeError(f"Termination unknown for {process.pid}; retain {root}") from error
            raise RuntimeError("Command exceeded deadline; no retry") from error
        calls.append({"binary": str(binary), "args": list(map(str, args)), "code": process.returncode,
                      "elapsed_seconds": round(time.monotonic() - began, 4),
                      "stdout": stdout, "stderr": stderr})
        (root / "commands.local.json").write_text(json.dumps(calls, indent=2))
        if process.returncode:
            raise RuntimeError(stderr or stdout)
        return json.loads(stdout)

    def approved(*args):
        preview = call(host, *args)
        return call(host, *args, "--approve", preview["approval_sha256"])

    spec = importlib.util.spec_from_file_location("log_report_native", EXAMPLE / "native.py")
    adapter = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(adapter)
    projects = adapter.prepare(root / "projects")
    runtime_file = root / "runtime.local.json"
    runtime = adapter.select_runtime(python, runtime_file)
    instances, releases = {}, {}
    for stage in ["workbook", "slides"]:
        project = projects / stage
        call(loom, "check", project)
        archive = call(loom, "pack", project)["archive"]
        call(loom, "verify", archive)
        delivery = root / (stage + "-release")
        releases[stage] = call(loom, "export-native-spool-v3", archive, delivery)["release"]
        instance = root / stage
        instances[stage] = instance
        loaded = call(host, "spool-load", delivery, instance, "--runtime", runtime_file)
        assert loaded["execution_granted"] is False and loaded["release"] == releases[stage]
        assert loaded["runtime_sha256"] == sha(runtime_file)
    a, b = instances["workbook"], instances["slides"]
    ws = root / "workspace"
    call(host, "spool-events-init", ws)
    route = root / "route.local.json"
    route.write_text(json.dumps({"schema": "an-native-artifact-route-v1", "source": str(a),
                                 "output": "report.xlsx", "target": str(b), "result": "report.pptx",
                                 "wall_seconds": 30}))
    approved("spool-event-subscribe", ws, "reports", route)
    seed = root / "requests.jsonl"
    shutil.copyfile(EXAMPLE / "sample.jsonl", seed)
    run = approved("spool-start", a, "first", seed, 30)
    assert run["status"] == "succeeded" and run["workspace"]["termination_confirmed"] is True
    accepted = approved("spool-accept", a, "first")
    old = accepted["result"]["outputs"][0]["binding"]["revision"]
    event = approved("spool-event-publish", ws, a, "first", "report.xlsx")["event"]
    book = root / "report.xlsx"
    first_artifacts = call(host, "spool-artifacts", a, "first")
    call(host, "artifact-checkout", first_artifacts["catalog"], old, book)
    # Create a newer upstream artifact before consuming the old event. This is
    # another real execution, never an automatic retry of an unknown outcome.
    with seed.open("a") as stream:
        stream.write('{"service":"api","status":503,"latency_ms":800}\n')
    newer = approved("spool-start", a, "newer", seed, 30)
    assert newer["status"] == "succeeded"
    accepted_new = approved("spool-accept", a, "newer")
    new = accepted_new["result"]["outputs"][0]["binding"]["revision"]
    assert old != new
    delivered = approved("spool-event-deliver", ws, event["id"], "reports")
    assert delivered["status"] == "succeeded" and delivered["consumed"] is False
    run = delivered["context"]["delivery"]
    assert delivered["result"]["workspace"]["termination_confirmed"] is True
    approved("spool-accept", b, run)
    complete = approved("spool-event-complete", ws, event["id"], "reports")
    assert complete["receipt"]["result_event"]["event"]["cause"] == delivered["context"]
    artifacts = call(host, "spool-artifacts", b, run)
    revision = artifacts["outputs"][0]["binding"]["revision"]
    deck = root / "report.pptx"
    call(host, "artifact-checkout", artifacts["catalog"], revision, deck)
    version = json.loads((Path(artifacts["catalog"]) / "revisions" / revision / "revision.json").read_text())
    assert version["source"]["input_revisions"][0]["revision"] == old
    # Domain validation uses the independently selected document libraries, not
    # a host/Loom special case or a screenshot of the result.
    inspect = """import json,sys
from openpyxl import load_workbook
from pptx import Presentation
w = load_workbook(sys.argv[1],read_only=True,data_only=False)
assert w['Report']['B8'].value == 4
w.close()
d = Presentation(sys.argv[2])
text = '\\n'.join(s.text for p in d.slides for s in p.shapes if s.has_text_frame)
assert '12 requests  /  3 errors' in text, text
assert len(d.slides) == 3
assert sum(s.has_table for p in d.slides for s in p.shapes) == 1
assert sum(s.shape_type == 13 for p in d.slides for s in p.shapes) == 0
print(json.dumps({'slides':len(d.slides),'native_tables':1,'raster_images':0}))
"""
    documents = call(python, "-B", "-s", "-c", inspect, book, deck)
    for instance, name in [(a, "first"), (a, "newer"), (b, run)]:
        approved("spool-reclaim", instance, name)
        approved("spool-reclaim-snapshots", instance, name)
    duplicate = call(host, "spool-event-deliver", ws, event["id"], "reports")
    assert duplicate["disposition"] == "already-consumed" and duplicate["executed"] is False
    retained = root / "retained.pptx"
    call(host, "artifact-checkout", artifacts["catalog"], revision, retained)
    assert retained.read_bytes() == deck.read_bytes()
    summary = {"schema": "loom-native-document-qualification-v1", "passed": True,
               "host_revision": options.host_revision, "host_sha256": sha(host), "loom_sha256": sha(loom),
               "releases": releases, "runtime_sha256": sha(runtime_file), "site_sha256": runtime["site_sha256"],
               "dependencies": runtime["dependencies"], "source_revision": old, "newer_revision": new,
               "result_revision": revision, "workbook_bytes": book.stat().st_size,
               "presentation_bytes": deck.stat().st_size, "documents": documents,
               "elapsed_seconds": round(time.monotonic() - started, 3),
               "model_invoked": False, "gui_edit_performed": False, "automatic_retries": 0}
    (root / "acceptance.json").write_text(json.dumps(summary, indent=2))
    print(json.dumps(summary))


if __name__ == "__main__":
    main()
