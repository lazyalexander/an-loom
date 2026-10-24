"""Domain acceptance remains outside Loom and an-agent implementations."""
import json
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "kit"))
import data
import workbook
import slides
from openpyxl import load_workbook
from pptx import Presentation

HERE = Path(__file__).resolve().parents[1]


class DomainChecks(unittest.TestCase):
    def test_metrics_and_saved_review_reach_native_editable_objects(self):
        with tempfile.TemporaryDirectory() as root:
            root = Path(root)
            book, deck = root / "report.xlsx", root / "report.pptx"
            workbook.create(HERE / "sample.jsonl", book)
            rows = workbook.read(book)
            self.assertEqual(sum(r["requests"] for r in rows), 12)
            self.assertEqual(sum(r["errors"] for r in rows), 3)
            self.assertEqual(rows[0]["p95_ms"], 650)
            w = load_workbook(book)
            w["Report"]["F8"] = "Investigate API timeouts"
            w.save(book)
            slides.create(book, deck)
            d = Presentation(deck)
            self.assertEqual(len(d.slides), 3)
            self.assertEqual(sum(s.has_table for p in d.slides for s in p.shapes), 1)
            self.assertIn("Investigate API timeouts", "\n".join(s.text for p in d.slides for s in p.shapes if s.has_text_frame))
            self.assertEqual(sum(s.shape_type == 13 for p in d.slides for s in p.shapes), 0)

    def test_unknown_logs_and_formula_cells_are_refused(self):
        with tempfile.TemporaryDirectory() as root:
            root = Path(root)
            source = root / "log.jsonl"
            source.write_text(json.dumps({"service":"api","status":200,"latency_ms":1,"command":"anything"}))
            with self.assertRaises(ValueError):
                data.aggregate(source)
            book = root / "report.xlsx"
            workbook.create(HERE / "sample.jsonl", book)
            w = load_workbook(book)
            w["Report"]["F8"] = '=HYPERLINK("https://example.invalid")'
            w.save(book)
            with self.assertRaises(ValueError):
                workbook.read(book)

    def test_external_workbook_parts_are_refused_before_library_loading(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "untrusted.xlsx"
            with zipfile.ZipFile(path, "w") as p:
                p.writestr("xl/externalLinks/externalLink1.xml", "external")
            with self.assertRaises(ValueError):
                workbook.read(path)


if __name__ == "__main__":
    unittest.main()
