"""Replaceable PPTX adapter: editable tables and text from validated report rows."""
import ooxml  # noqa: F401; initialize deterministic MIME before document libraries
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from workbook import read

INK = RGBColor.from_string("17344A")


def text(slide, value, x, y, width, height, size=22, bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    shape.text = value
    for paragraph in shape.text_frame.paragraphs:
        paragraph.font.name = "Calibri"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = INK
    return shape


def create(source, destination):
    rows = read(source)
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.33), Inches(7.5)
    cover = deck.slides.add_slide(deck.slide_layouts[6])
    text(cover, "Service reliability", 0.8, 1.25, 11.7, 1, 44, True)
    text(cover, "Synthetic operations report", 0.85, 2.45, 11, .6, 24)
    text(cover, f'{sum(r["requests"] for r in rows)} requests  /  {sum(r["errors"] for r in rows)} errors',
         .85, 4.2, 11.5, 1, 34)
    text(cover, "Source: saved report workbook", .85, 6.5, 11, .4, 17)
    metrics = deck.slides.add_slide(deck.slide_layouts[6])
    text(metrics, "Service metrics", .75, .45, 11.8, .7, 34, True)
    table = metrics.shapes.add_table(len(rows) + 1, 5, Inches(.8), Inches(1.65),
                                     Inches(11.7), Inches(3.7)).table
    headers = ["Service", "Requests", "Errors", "p95 (ms)", "Owner"]
    for i, width in enumerate([2.5, 2, 1.7, 2, 3.5]):
        table.columns[i].width = Inches(width)
    for i, row in enumerate([headers] + [[r["service"], str(r["requests"]), str(r["errors"]),
                                         str(r["p95_ms"]), r["owner"] or "Unassigned"] for r in rows]):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = INK if i == 0 else RGBColor.from_string("EDF3F7")
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = "Calibri"
                paragraph.font.size = Pt(20)
                paragraph.font.bold = i == 0
                paragraph.font.color.rgb = RGBColor(255, 255, 255) if i == 0 else INK
    actions = deck.slides.add_slide(deck.slide_layouts[6])
    text(actions, "Review actions", .75, .45, 11.8, .7, 34, True)
    for i, row in enumerate(rows):
        # This intentionally small report supports at most three services in the
        # slide layout. The domain adapter refuses overflow rather than hiding it.
        if i >= 3:
            raise ValueError("This slide layout supports at most three services")
        text(actions, row["service"], .85, 1.55 + i * 1.65, 3, .5, 24, True)
        text(actions, row["note"] or "No review note recorded", 4, 1.55 + i * 1.65, 8.3, 1.25, 22)
    deck.save(destination)
