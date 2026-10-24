"""Replaceable XLSX adapter; no host, engine or model imports."""
import ooxml  # noqa: F401; initialize deterministic MIME before document libraries
import hashlib
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from data import aggregate, check_package, validate_rows

HEADER = ["Service", "Requests", "Errors", "p95 latency (ms)", "Owner", "Review note"]


def create(source, destination):
    rows = aggregate(source)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:F1")
    sheet["A1"] = "Service reliability"
    sheet["A1"].font = Font(name="Calibri", size=24, bold=True, color="17344A")
    sheet["A3"], sheet["C3"], sheet["E3"] = "Requests", "Errors", "Error rate"
    last = 7 + len(rows)
    sheet["B3"], sheet["D3"], sheet["F3"] = f"=SUM(B8:B{last})", f"=SUM(C8:C{last})", "=IF(B3=0,0,D3/B3)"
    sheet["F3"].number_format = "0.0%"
    sheet.merge_cells("A5:F5")
    sheet["A5"] = "Synthetic logs. Edit Owner and Review note before saving a new report version."
    sheet["A5"].font = Font(name="Calibri", size=11, color="596B79")
    for col, value in enumerate(HEADER, 1):
        cell = sheet.cell(7, col, value)
        cell.fill = PatternFill("solid", fgColor="17344A")
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    for index, row in enumerate(rows, 8):
        for col, key in enumerate(["service", "requests", "errors", "p95_ms", "owner", "note"], 1):
            cell = sheet.cell(index, col, row[key])
            cell.font = Font(name="Calibri", size=12, color="1769AA" if col >= 5 else "243746")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if col >= 5:
                cell.fill = PatternFill("solid", fgColor="EAF4FA")
        sheet.row_dimensions[index].height = 42
    for col, width in {"A": 20, "B": 14, "C": 12, "D": 21, "E": 23, "F": 52}.items():
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = "B8"
    sheet.auto_filter.ref = f"A7:F{last}"
    sheet.print_options.horizontalCentered = True
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.print_area = f"A1:F{last}"
    provenance = workbook.create_sheet("Source")
    provenance.append(["Input SHA-256", hashlib.sha256(Path(source).read_bytes()).hexdigest()])
    provenance.append(["Method", "Nearest-rank p95; errors are HTTP status >= 500"])
    provenance.column_dimensions["A"].width = 23
    provenance.column_dimensions["B"].width = 75
    workbook.save(destination)


def read(source):
    check_package(source)
    # Require defusedxml before opening user-edited OOXML. Do not evaluate macros,
    # external links or arbitrary formulas. Summary formulas are display-only.
    import defusedxml  # noqa: F401
    workbook = load_workbook(source, read_only=True, data_only=False, keep_links=False)
    try:
        sheet = workbook["Report"]
        if sheet.max_row > 27 or sheet.max_column > 6:
            raise ValueError("Unexpected workbook dimensions")
        if [sheet.cell(7, c).value for c in range(1, 7)] != HEADER:
            raise ValueError("Unsupported report columns")
        rows = []
        for values in sheet.iter_rows(min_row=8, max_col=6, values_only=True):
            if all(v is None for v in values):
                continue
            service, requests, errors, latency, owner, note = values
            rows.append(dict(service=service, requests=requests, errors=errors, p95_ms=latency,
                             owner=owner or "", note=note or ""))
        return validate_rows(rows)
    finally:
        workbook.close()
